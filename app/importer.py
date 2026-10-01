import csv
import io
import json
import re
from datetime import datetime, timezone

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

SYSTEM_FIELDS = [
    {"key": "code", "label": "Buyurtma kodi", "required": True},
    {"key": "client_name", "label": "Mijoz nomi", "required": True},
    {"key": "pickup_address", "label": "Olish manzili / ombor", "required": False},
    {"key": "dropoff_address", "label": "Yetkazish manzili", "required": False},
    {"key": "cargo", "label": "Yuk / tur", "required": False},
    {"key": "weight_kg", "label": "Og‘irlik (kg)", "required": False},
    {"key": "amount", "label": "Summa", "required": False},
    {"key": "route_code", "label": "Yuklama / marshrut №", "required": False},
    {"key": "dropoff_lat", "label": "Yetkazish latitude (Y)", "required": False},
    {"key": "dropoff_lng", "label": "Yetkazish longitude (X)", "required": False},
    {"key": "notes", "label": "Izoh", "required": False},
    {"key": "client_code", "label": "Mijoz kodi", "required": False},
    {"key": "sales_rep", "label": "Agent (savdo vakili)", "required": False},
    {"key": "agent_code", "label": "Kod torгового", "required": False},
    {"key": "driver_name", "label": "Haydovchi", "required": False},
    {"key": "vehicle_plate", "label": "Davlat raqami", "required": False},
    {"key": "delivery_date", "label": "Yetkazish sanasi", "required": False},
    {"key": "payment_status", "label": "To‘lov holati", "required": False},
]

DRIVER_FIELDS = [
    {"key": "name", "label": "Ism", "required": True},
    {"key": "phone", "label": "Telefon", "required": False},
    {"key": "vehicle_plate", "label": "Davlat raqami", "required": False},
    {"key": "vehicle_type", "label": "Transport turi", "required": False},
    {"key": "status", "label": "Holat", "required": False},
    {"key": "agent_name", "label": "Agent", "required": False},
]

DRIVER_MAPPING = {
    "name": "Ism",
    "phone": "Telefon",
    "vehicle_plate": "Davlat raqami",
    "vehicle_type": "Transport turi",
    "status": "Holat",
    "agent_name": "Agent",
}

DRIVER_ALIASES = {
    "name": ["ism", "fio", "фио", "haydovchi", "driver", "name", "имя", "ф.и.о"],
    "phone": ["telefon", "phone", "тел", "телефон"],
    "vehicle_plate": ["davlat raqami", "raqam", "plate", "гос номер", "госномер", "номер авто"],
    "vehicle_type": ["transport turi", "turi", "type", "тип", "mashina"],
    "status": ["holat", "status", "статус"],
    "agent_name": ["agent", "агент", "savdo", "agent kod", "kod"],
}

ORDER_ALIASES = {
    "code": ["id", "код", "orderid", "order id", "buyurtma", "номер"],
    "client_name": ["имя клиен", "клиент", "mijoz", "customer", "client"],
    "pickup_address": ["название склада", "склад", "ombor", "from", "pickup"],
    "dropoff_address": ["имя филиала", "филиал", "манзил", "to", "dropoff", "адрес"],
    "cargo": ["тип", "yuk", "goods", "товар"],
    "weight_kg": ["кг", "weight", "ogirlik", "вес"],
    "amount": ["сумма", "summa", "narx"],
    "route_code": ["загрузка", "yuklama", "маршрут"],
    "dropoff_lat": ["gpscoordy", "lat", "latitude", "y"],
    "dropoff_lng": ["gpscoordx", "lng", "longitude", "x"],
    "notes": ["примеч", "izoh", "note"],
    "client_code": ["код клиен", "mijoz kodi"],
    "sales_rep": ["имя торгового", "agent", "торговый", "savdo"],
    "agent_code": [
        "код торгового",
        "код торгов",
        "kod torgovogo",
        "код агента",
        "agent kod",
        "код торг",
    ],
    "driver_name": [
        "водитель",
        "имя водителя",
        "фио водителя",
        "экспедитор",
        "haydovchi",
        "курьер",
        "доставщик",
        "courier",
        "driver",
    ],
    "vehicle_plate": [
        "гос номер",
        "госномер",
        "гос. номер",
        "номер авто",
        "номер машины",
        "davlat raqami",
        "plate",
        "транспорт",
    ],
    "delivery_date": ["дата доставки", "yetkazish", "delivery"],
    "payment_status": ["статус платежа", "tolov"],
}


def fields_for(entity: str) -> list[dict]:
    if (entity or "orders") == "drivers":
        return DRIVER_FIELDS
    return SYSTEM_FIELDS


def build_excel_bytes(headers: list[str], sheet_name: str = "Shablon", sample_row: list[str] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = (sheet_name or "Shablon")[:31]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F3A5F")
    ws.append(headers)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
    if sample_row:
        ws.append(sample_row)
    for idx, header in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(idx)].width = max(16, min(36, len(str(header)) + 8))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def headers_from_mapping(mapping: dict[str, str], entity: str = "orders") -> list[str]:
    fields = fields_for(entity)
    headers = []
    seen = set()
    for field in fields:
        col = mapping.get(field["key"]) or field["label"]
        if col and col not in seen:
            headers.append(col)
            seen.add(col)
    for col in mapping.values():
        if col and col not in seen:
            headers.append(col)
            seen.add(col)
    return headers or [f["label"] for f in fields]


def normalize_header(value) -> str:
    return " ".join(str(value or "").replace("\n", " ").replace("\r", " ").split()).strip()


def header_key(value) -> str:
    return normalize_header(value).lower()


def loose_key(value) -> str:
    return re.sub(r"[^\wа-яё]+", " ", header_key(value), flags=re.IGNORECASE).strip()


def parse_template_config(mapping_json: str) -> tuple[str | int | None, int, dict[str, str]]:
    data = json.loads(mapping_json) if mapping_json else {}
    if "mapping" in data and isinstance(data["mapping"], dict):
        return data.get("sheet") or None, int(data.get("header_row") or 1), data["mapping"]
    meta_keys = {"sheet", "header_row", "mapping"}
    mapping = {k: v for k, v in data.items() if k not in meta_keys}
    return data.get("sheet"), int(data.get("header_row") or 1), mapping


def build_template_json(sheet, header_row: int, mapping: dict[str, str]) -> str:
    return json.dumps(
        {"sheet": sheet, "header_row": header_row, "mapping": mapping},
        ensure_ascii=False,
    )


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value).strip()
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


DRIVER_COL_ORDER = ["name", "phone", "vehicle_plate", "vehicle_type", "status", "agent_name"]
HEADER_LIKE_NAMES = {
    "ism",
    "fio",
    "name",
    "haydovchi",
    "driver",
    "имя",
    "фио",
    "ф.и.о",
}


def _driver_header_wanted() -> set[str]:
    wanted = {header_key(v) for v in DRIVER_MAPPING.values()}
    for aliases in DRIVER_ALIASES.values():
        wanted.update(header_key(a) for a in aliases)
    return wanted


def _header_matches_driver(headers: list[str]) -> bool:
    wanted = _driver_header_wanted()
    return any(header_key(h) in wanted for h in headers if h)


def record_cells(record: dict) -> list[str]:
    raw = record.get("_cells")
    if isinstance(raw, list):
        return [_cell(x) for x in raw]
    return [_cell(v) for k, v in record.items() if k != "_cells"]


def driver_value(record: dict, mapping: dict[str, str], field: str) -> str:
    clean = {k: v for k, v in record.items() if k != "_cells"}
    val = mapped_value(clean, mapping, field)
    if val:
        return val[:200]
    cells = record_cells(record)
    try:
        idx = DRIVER_COL_ORDER.index(field)
    except ValueError:
        return ""
    if idx < len(cells) and cells[idx]:
        return cells[idx][:200]
    return ""


def _rows_to_driver_records(rows: list[list], header_row: int = 1) -> list[dict[str, str]]:
    if not rows:
        return []
    idx = max(int(header_row or 1) - 1, 0)
    if idx >= len(rows):
        return []
    headers = [normalize_header(h) for h in rows[idx]]
    start = idx + 1
    std_headers = list(DRIVER_MAPPING.values())
    if not _header_matches_driver(headers):
        headers = std_headers
        start = idx
    records: list[dict[str, str]] = []
    for row in rows[start:]:
        cells = [_cell(x) for x in row]
        while cells and not cells[-1]:
            cells.pop()
        if not any(cells):
            continue
        rec: dict = {"_cells": cells}
        for i, h in enumerate(headers):
            if not h:
                continue
            rec[h] = cells[i] if i < len(cells) else ""
        for i, h in enumerate(std_headers):
            if h not in rec and i < len(cells):
                rec[h] = cells[i]
        if any(str(v).strip() for k, v in rec.items() if k != "_cells"):
            records.append(rec)
    return records


def read_driver_records(raw: bytes, filename: str, header_row: int = 1) -> list[dict[str, str]]:
    name = (filename or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        wb = load_workbook(io.BytesIO(raw), data_only=False, read_only=False)
        ws = wb[wb.sheetnames[0]]
        rows = [list(row) for row in ws.iter_rows(values_only=True)]
        wb.close()
        return _rows_to_driver_records(rows, header_row)
    text = raw.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    return _rows_to_driver_records(rows, header_row)


def inspect_table(raw: bytes, filename: str) -> dict:
    name = (filename or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        return _inspect_xlsx(raw)
    return _inspect_csv(raw)


def _inspect_csv(raw: bytes) -> dict:
    text = raw.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    rows = [list(r) for r in reader]
    headers = [normalize_header(h) for h in (rows[0] if rows else [])]
    sample = []
    for row in rows[1:6]:
        sample.append({headers[i]: _cell(row[i]) if i < len(row) else "" for i in range(len(headers))})
    return {"sheets": [{"name": "CSV", "headers": headers, "sample": sample, "row_count": max(len(rows) - 1, 0)}]}


def _inspect_xlsx(raw: bytes) -> dict:
    wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    sheets = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            rows.append(list(row))
            if i >= 80:
                break
        headers = [normalize_header(h) for h in (rows[0] if rows else []) if normalize_header(h)]
        if rows:
            raw_headers = [normalize_header(h) for h in rows[0]]
            headers = raw_headers
        sample = []
        for row in rows[1:6]:
            item = {}
            for i, h in enumerate(headers):
                if not h:
                    continue
                item[h] = _cell(row[i]) if i < len(row) else ""
            sample.append(item)
        sheets.append(
            {
                "name": sheet_name,
                "headers": [h for h in headers if h],
                "sample": sample,
                "row_count": None,
            }
        )
    wb.close()
    return {"sheets": sheets}


def read_records(raw: bytes, filename: str, sheet, header_row: int) -> list[dict[str, str]]:
    name = (filename or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        return _read_xlsx(raw, sheet, header_row)
    return _read_csv(raw, header_row)


def _read_csv(raw: bytes, header_row: int) -> list[dict[str, str]]:
    text = raw.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    idx = max(header_row - 1, 0)
    if idx >= len(rows):
        return []
    return _rows_to_order_records(rows, idx)


def _rows_to_order_records(rows: list[list], header_idx: int) -> list[dict[str, str]]:
    if header_idx >= len(rows):
        return []
    headers = [normalize_header(h) for h in rows[header_idx]]
    records = []
    for row in rows[header_idx + 1 :]:
        rec = {}
        for i, h in enumerate(headers):
            if not h:
                continue
            rec[h] = _cell(row[i]) if i < len(row) else ""
        if any(str(v).strip() for v in rec.values()):
            records.append(rec)
    return records


def _open_xlsx(raw: bytes):
    try:
        return load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    except Exception:
        return load_workbook(io.BytesIO(raw), data_only=False, read_only=False)


def _read_xlsx(raw: bytes, sheet, header_row: int) -> list[dict[str, str]]:
    wb = _open_xlsx(raw)
    try:
        if sheet and str(sheet) in wb.sheetnames:
            ws = wb[str(sheet)]
        elif isinstance(sheet, int) and 0 <= sheet < len(wb.sheetnames):
            ws = wb[wb.sheetnames[sheet]]
        else:
            ws = wb[wb.sheetnames[0]]
        rows = [list(row) for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    idx = max(int(header_row or 1) - 1, 0)
    if idx >= len(rows):
        return []
    return _rows_to_order_records(rows, idx)


def mapped_value(record: dict[str, str], mapping: dict[str, str], field: str) -> str:
    column = mapping.get(field) or ""
    aliases = list(DRIVER_ALIASES.get(field, []))
    aliases.extend(ORDER_ALIASES.get(field, []))
    wanted = {header_key(x) for x in [column, field, *aliases] if x}
    wanted_loose = {loose_key(x) for x in [column, field, *aliases] if x}
    items = [(k, v) for k, v in record.items() if k != "_cells"]
    lookup = dict(items)
    if column in lookup:
        val = (lookup.get(column) or "").strip()
        if val:
            return val
    for key, value in items:
        if header_key(key) in wanted or loose_key(key) in wanted_loose:
            val = (value or "").strip()
            if val:
                return val
    return ""


def to_float(value: str, default: float = 0.0) -> float:
    if not value:
        return default
    text = str(value).replace(" ", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return default


RELOG_MAPPING = {
    "code": "ID",
    "client_name": "Имя Клиен.",
    "pickup_address": "Название склада",
    "dropoff_address": "Имя филиала",
    "cargo": "Тип",
    "weight_kg": "кг",
    "amount": "Сумма",
    "route_code": "Загрузка №",
    "dropoff_lat": "GpsCoordY",
    "dropoff_lng": "GpsCoordX",
    "notes": "Примеч.",
    "client_code": "Код клиен.",
    "sales_rep": "Имя торгового",
    "agent_code": "Код торгового",
    "driver_name": "Водитель",
    "vehicle_plate": "Гос. номер",
    "delivery_date": "Дата доставки",
    "payment_status": "Статус платежа",
}

LABEL_MAPPING = {f["key"]: f["label"] for f in SYSTEM_FIELDS}


def mapping_score(headers: list[str], mapping: dict[str, str]) -> int:
    found = {loose_key(h) for h in headers if h}
    score = 0
    for field, col in mapping.items():
        aliases = [col, field, *ORDER_ALIASES.get(field, [])]
        if any(loose_key(a) in found for a in aliases if a):
            score += 2 if field in {"code", "client_name", "dropoff_address"} else 1
    return score


def looks_like_xlsx(raw: bytes, filename: str) -> bool:
    name = (filename or "").lower()
    if raw[:2] == b"PK":
        return True
    return name.endswith((".xlsx", ".xlsm", ".xltx"))


def looks_like_xls(raw: bytes, filename: str) -> bool:
    name = (filename or "").lower()
    return raw[:4] == b"\xd0\xcf\x11\xe0" or name.endswith(".xls")


def _sheet_rows(raw: bytes, filename: str) -> list[tuple[str, list[list]]]:
    if looks_like_xls(raw, filename) and not looks_like_xlsx(raw, filename):
        raise ValueError("Eski .xls fayl. Excel’da «Saqlash .xlsx» qilib qayta yuklang.")
    if looks_like_xlsx(raw, filename):
        wb = _open_xlsx(raw)
        try:
            out = []
            for name in wb.sheetnames:
                ws = wb[name]
                rows = [list(row) for row in ws.iter_rows(values_only=True)]
                out.append((name, rows))
            return out
        finally:
            wb.close()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1251", errors="replace")
    return [("CSV", list(csv.reader(io.StringIO(text))))]


def auto_read_order_records(
    raw: bytes,
    filename: str,
    mappings: list[dict[str, str]] | None = None,
) -> tuple[list[dict[str, str]], dict[str, str], dict]:
    candidates = [m for m in (mappings or []) if m]
    if RELOG_MAPPING not in candidates:
        candidates.append(RELOG_MAPPING)
    if LABEL_MAPPING not in candidates:
        candidates.append(LABEL_MAPPING)
    tables = _sheet_rows(raw, filename)
    best = None
    for sheet_name, rows in tables:
        if not rows:
            continue
        limit = min(len(rows), 25)
        for idx in range(limit):
            headers = [normalize_header(h) for h in rows[idx]]
            if sum(1 for h in headers if h) < 3:
                continue
            for mapping in candidates:
                score = mapping_score(headers, mapping)
                if best is None or score > best["score"]:
                    best = {
                        "score": score,
                        "sheet": sheet_name,
                        "header_idx": idx,
                        "mapping": mapping,
                        "rows": rows,
                    }
    if not best:
        raise ValueError("Jadvalda ustunlar topilmadi")
    records = _rows_to_order_records(best["rows"], best["header_idx"])
    return records, best["mapping"], {"sheet": best["sheet"], "header_row": best["header_idx"] + 1, "score": best["score"]}
