"""Pivot-hisobot kutubxonasi: zayavkalarni istalgan maydonlar bo‘yicha guruhlaydi (qatorlar × ustunlar × qiymatlar).

Maydon qiymatlari matn ko‘rinishida olinadi, shuning uchun qiymat filtrlari va saqlangan tuzilmalar
tashkilot ichidagi ID’larga bog‘lanib qolmaydi.
"""

from __future__ import annotations

import io
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .dispatch import LOCAL_TZ
from .reports import PROOF_LABELS, STATUS_LABELS, OrderRow, _clean

EMPTY = "(bo‘sh)"
OTHERS = "Boshqalar"
WEEKDAYS = ["Dushanba", "Seshanba", "Chorshanba", "Payshanba", "Juma", "Shanba", "Yakshanba"]
MAX_ROWS_DIMS = 8
MAX_COLS_DIMS = 2
MAX_VALUES = 12
MAX_FILTER_FIELDS = 20
MAX_FILTER_VALUES = 1000


def _txt(value) -> str:
    text = str(value or "").strip()
    return text or EMPTY


def _join(*parts) -> str:
    return _txt(" · ".join(str(p).strip() for p in parts if str(p or "").strip()))


def _week(day: str) -> str:
    try:
        d = date.fromisoformat(day)
    except (TypeError, ValueError):
        return EMPTY
    start = d - timedelta(days=d.weekday())
    year, week, _ = d.isocalendar()
    return f"{year}-H{week:02d} ({start:%d.%m}–{start + timedelta(days=6):%d.%m})"


def _weekday(day: str) -> str:
    try:
        d = date.fromisoformat(day)
    except (TypeError, ValueError):
        return EMPTY
    return f"{d.weekday() + 1}. {WEEKDAYS[d.weekday()]}"


def _proof(r: OrderRow) -> str:
    if r.status == "returned":
        return "Qaytarildi"
    return PROOF_LABELS.get(r.proof_reason) or ("Tasdiqsiz yetkazildi" if r.status == "delivered" else "Tasdiq yo‘q")


@dataclass(frozen=True)
class Dim:
    key: str
    label: str
    group: str
    get: callable


DIMENSIONS: dict[str, Dim] = {
    d.key: d
    for d in [
        Dim("org", "Tashkilot", "Asosiy", lambda r: _txt(r.org_name)),
        Dim("status", "Holat", "Asosiy", lambda r: _txt(STATUS_LABELS.get(r.status, r.status))),
        Dim("agent", "Agent", "Asosiy", lambda r: _join(r.agent_code, r.agent_name)),
        Dim("driver", "Haydovchi", "Asosiy", lambda r: _txt(r.driver_name)),
        Dim("client", "Klient", "Asosiy", lambda r: _join(r.client_code, r.client_name)),
        Dim("warehouse", "Sklad", "Asosiy", lambda r: _txt(r.warehouse_name)),
        Dim("payment", "To‘lov holati", "Asosiy", lambda r: _txt(r.payment)),
        Dim("proof", "Yetkazish belgisi", "Tasdiq", _proof),
        Dim("comment", "Haydovchi izohi", "Tasdiq", lambda r: _txt(r.proof_comment)),
        Dim("has_photo", "Rasm bormi", "Tasdiq", lambda r: "Rasm bor" if r.proof_photo else "Rasm yo‘q"),
        Dim("day", "Sana", "Vaqt", lambda r: _txt(r.day)),
        Dim("week", "Hafta", "Vaqt", lambda r: _week(r.day)),
        Dim("month", "Oy", "Vaqt", lambda r: _txt(r.day[:7])),
        Dim("weekday", "Hafta kuni", "Vaqt", lambda r: _weekday(r.day)),
        Dim("proof_hour", "Tasdiq soati", "Vaqt", lambda r: f"{r.proof_at[11:13]}:00" if len(r.proof_at) >= 13 else EMPTY),
        Dim("cargo", "Savdo nuqtasi turi", "Qo‘shimcha", lambda r: _txt(r.cargo)),
        Dim("route", "Yo‘nalish kodi", "Qo‘shimcha", lambda r: _txt(r.route_code)),
        Dim("address", "Manzil", "Qo‘shimcha", lambda r: _txt(r.address)),
        Dim("order", "Zayavka (har biri alohida)", "Qo‘shimcha", lambda r: _txt(r.code)),
    ]
}


@dataclass(frozen=True)
class Measure:
    key: str
    label: str
    kind: str  # int | money | float | pct


MEASURES: dict[str, Measure] = {
    m.key: m
    for m in [
        Measure("count", "Zayavkalar soni", "int"),
        Measure("amount", "Savdo summasi", "money"),
        Measure("delivered", "Yetkazildi", "int"),
        Measure("returned", "Qaytarildi", "int"),
        Measure("pending", "Jarayonda", "int"),
        Measure("amount_delivered", "Yetkazilgan summa", "money"),
        Measure("amount_returned", "Qaytarilgan summa", "money"),
        Measure("avg_check", "O‘rtacha chek", "money"),
        Measure("return_rate", "Qaytarish %", "pct"),
        Measure("weight", "Og‘irlik, kg", "float"),
        Measure("clients", "Klientlar soni", "int"),
        Measure("drivers", "Haydovchilar soni", "int"),
    ]
}


def catalog() -> dict:
    return {
        "dimensions": [{"key": d.key, "label": d.label, "group": d.group} for d in DIMENSIONS.values()],
        "measures": [{"key": m.key, "label": m.label, "kind": m.kind} for m in MEASURES.values()],
    }


class Acc:
    __slots__ = ("count", "amount", "amount_delivered", "amount_returned", "weight", "delivered", "returned", "pending", "client_set", "driver_set")

    def __init__(self) -> None:
        self.count = 0
        self.amount = 0.0
        self.amount_delivered = 0.0
        self.amount_returned = 0.0
        self.weight = 0.0
        self.delivered = 0
        self.returned = 0
        self.pending = 0
        self.client_set: set = set()
        self.driver_set: set = set()

    def add(self, r: OrderRow) -> None:
        self.count += 1
        self.amount += r.amount
        self.weight += r.weight
        if r.status == "delivered":
            self.delivered += 1
            self.amount_delivered += r.amount
        elif r.status == "returned":
            self.returned += 1
            self.amount_returned += r.amount
        else:
            self.pending += 1
        self.client_set.add(r.client_key)
        if r.driver_id:
            self.driver_set.add(r.driver_id)

    def value(self, key: str):
        if key == "avg_check":
            return round(self.amount / self.count, 2) if self.count else 0
        if key == "return_rate":
            done = self.delivered + self.returned
            return round(self.returned * 100 / done, 1) if done else 0
        if key == "clients":
            return len(self.client_set)
        if key == "drivers":
            return len(self.driver_set)
        v = getattr(self, key)
        return round(v, 2) if isinstance(v, float) else v

    def values(self, keys: list[str]) -> list:
        return [self.value(k) for k in keys]


@dataclass
class PivotConfig:
    rows: list[str] = field(default_factory=list)
    cols: list[str] = field(default_factory=list)
    values: list[str] = field(default_factory=lambda: ["count"])
    filters: dict[str, list[str]] = field(default_factory=dict)
    subtotals: bool = False
    sort_by: str = ""
    sort_dir: str = "desc"

    def validate(self) -> None:
        for name, keys, limit in (("Qatorlar", self.rows, MAX_ROWS_DIMS), ("Ustunlar", self.cols, MAX_COLS_DIMS)):
            if len(keys) > limit:
                raise ValueError(f"{name}: ko‘pi bilan {limit} ta maydon")
            bad = [k for k in keys if k not in DIMENSIONS]
            if bad:
                raise ValueError(f"Noma’lum maydon: {', '.join(bad)}")
        if set(self.rows) & set(self.cols):
            raise ValueError("Bir maydon ham qatorda, ham ustunda bo‘la olmaydi")
        if len(set(self.rows)) != len(self.rows) or len(set(self.cols)) != len(self.cols):
            raise ValueError("Maydon takrorlangan")
        self.values = [v for v in dict.fromkeys(self.values) if v in MEASURES][:MAX_VALUES] or ["count"]
        if len(self.filters) > MAX_FILTER_FIELDS:
            raise ValueError(f"Filtr: ko‘pi bilan {MAX_FILTER_FIELDS} ta maydon")
        for k, vals in self.filters.items():
            if k not in DIMENSIONS:
                raise ValueError(f"Noma’lum filtr maydoni: {k}")
            if len(vals) > MAX_FILTER_VALUES:
                raise ValueError(f"«{DIMENSIONS[k].label}» filtrida juda ko‘p qiymat tanlangan")
        if self.sort_by and self.sort_by not in MEASURES:
            self.sort_by = ""
        if self.sort_dir not in ("asc", "desc"):
            self.sort_dir = "desc"

    def describe(self) -> list[tuple[str, str]]:
        out = [
            ("Qatorlar", ", ".join(DIMENSIONS[k].label for k in self.rows) or "—"),
            ("Ustunlar", ", ".join(DIMENSIONS[k].label for k in self.cols) or "—"),
            ("Qiymatlar", ", ".join(MEASURES[k].label for k in self.values)),
        ]
        for k, vals in self.filters.items():
            if vals:
                shown = ", ".join(vals[:8]) + (f" (+{len(vals) - 8})" if len(vals) > 8 else "")
                out.append((f"Filtr: {DIMENSIONS[k].label}", shown))
        return out


def apply_value_filters(rows: list[OrderRow], filters: dict[str, list[str]], skip: str = "") -> list[OrderRow]:
    checks = [(DIMENSIONS[k].get, set(v)) for k, v in filters.items() if v and k != skip and k in DIMENSIONS]
    if not checks:
        return rows
    return [r for r in rows if all(fn(r) in vals for fn, vals in checks)]


def distinct_values(rows: list[OrderRow], key: str, limit: int = 2000) -> list[dict]:
    fn = DIMENSIONS[key].get
    counts = Counter(fn(r) for r in rows)
    items = sorted(counts.items(), key=lambda kv: (-kv[1], kv[0].lower()))[:limit]
    return [{"value": v, "count": n} for v, n in items]


STATUS_RANK = {label: i for i, label in enumerate(STATUS_LABELS.values())}


def _key_sort(keys: tuple) -> tuple:
    return tuple((k == EMPTY, k == OTHERS, STATUS_RANK.get(k, 99), k.lower()) for k in keys)


def run_pivot(data: list[OrderRow], cfg: PivotConfig, row_limit: int = 5000, col_limit: int = 60) -> dict:
    rows = apply_value_filters(data, cfg.filters)
    rdims = [DIMENSIONS[k].get for k in cfg.rows]
    cdims = [DIMENSIONS[k].get for k in cfg.cols]
    measures = cfg.values

    col_key = (lambda r: tuple(fn(r) for fn in cdims)) if cdims else (lambda r: ())
    allowed = None
    if cdims:
        freq = Counter(col_key(r) for r in rows)
        if len(freq) > col_limit:
            allowed = {k for k, _ in freq.most_common(col_limit - 1)}

    cells: dict = defaultdict(Acc)
    row_tot: dict = defaultdict(Acc)
    col_tot: dict = defaultdict(Acc)
    # Oraliq jamilar: har bir prefiks (masalan tashkilot, tashkilot+agent) bo‘yicha, oxirgi darajadan tashqari
    levels = len(rdims) - 1 if cfg.subtotals and len(rdims) > 1 else 0
    sub_tot: dict = defaultdict(Acc)
    grand = Acc()
    for r in rows:
        rk = tuple(fn(r) for fn in rdims)
        ck = col_key(r)
        if allowed is not None and ck not in allowed:
            ck = (OTHERS,) + ("",) * (len(cdims) - 1)
        grand.add(r)
        row_tot[rk].add(r)
        if cdims:
            cells[(rk, ck)].add(r)
            col_tot[ck].add(r)
        for lvl in range(levels):
            prefix = rk[: lvl + 1]
            sub_tot[prefix].add(r)
            if cdims:
                cells[(prefix, ck)].add(r)

    col_keys = sorted(col_tot, key=_key_sort)
    sort_m = cfg.sort_by
    sign = -1 if cfg.sort_dir == "desc" else 1
    keys = list(row_tot)
    if sort_m:
        def sort_key(k):
            parts = [(sign * sub_tot[k[: lvl + 1]].value(sort_m), _key_sort((k[lvl],))) for lvl in range(levels)]
            parts.append((sign * row_tot[k].value(sort_m), _key_sort(k)))
            return parts

        keys.sort(key=sort_key)
    else:
        keys.sort(key=_key_sort)

    empty = [None] * len(measures)

    def cell_values(key: tuple) -> list[list]:
        if not cdims:
            return []
        return [cells[(key, ck)].values(measures) if (key, ck) in cells else empty for ck in col_keys]

    out_rows: list[dict] = []
    truncated = False
    detail_rows: Counter = Counter()
    for k in row_tot:
        for lvl in range(levels):
            detail_rows[k[: lvl + 1]] += 1

    def close_groups(prev: tuple, upto: int) -> None:
        for lvl in range(levels - 1, upto - 1, -1):
            prefix = prev[: lvl + 1]
            if detail_rows[prefix] < 2:
                continue
            out_rows.append({"type": "subtotal", "level": lvl, "keys": list(prefix), "cells": cell_values(prefix), "total": sub_tot[prefix].values(measures)})

    prev = None
    for n, rk in enumerate(keys):
        if n >= row_limit:
            truncated = True
            break
        if prev is not None and levels:
            diff = next((lvl for lvl in range(levels) if rk[lvl] != prev[lvl]), None)
            if diff is not None:
                close_groups(prev, diff)
        out_rows.append({"type": "row", "keys": list(rk), "cells": cell_values(rk), "total": row_tot[rk].values(measures)})
        prev = rk
    if prev is not None and levels and not truncated:
        close_groups(prev, 0)

    return {
        "row_fields": [{"key": k, "label": DIMENSIONS[k].label} for k in cfg.rows],
        "col_fields": [{"key": k, "label": DIMENSIONS[k].label} for k in cfg.cols],
        "values": [{"key": k, "label": MEASURES[k].label, "kind": MEASURES[k].kind} for k in measures],
        "col_keys": [list(k) for k in col_keys],
        "rows": out_rows,
        "total": {"cells": [col_tot[ck].values(measures) for ck in col_keys], "total": grand.values(measures)},
        "row_count": len(keys),
        "truncated": truncated,
        "orders_used": len(rows),
        "orders_total": len(data),
        "cols_merged": allowed is not None,
    }


# ---------- Excel ----------

HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
COL_FILL = PatternFill("solid", fgColor="2D5A8C")
SUB_FILL = PatternFill("solid", fgColor="E8EEF8")
TOTAL_FILL = PatternFill("solid", fgColor="D3DEEF")
BOLD = Font(bold=True)
FORMATS = {"money": "#,##0", "int": "#,##0", "float": "#,##0.0", "pct": "0.0"}


def _cell(ws, value, *, font=None, fill=None, fmt=None, wrap=False):
    c = WriteOnlyCell(ws, value=_clean(value))
    if font:
        c.font = font
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if wrap:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    return c


def build_pivot_xlsx(
    result: dict,
    cfg: PivotConfig,
    title: str,
    meta: list[tuple[str, str]],
    raw_rows: list[OrderRow],
    generated_by: str,
    raw_cap: int = 200_000,
) -> bytes:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Pivot")
    dims = result["row_fields"]
    vals = result["values"]
    col_keys = result["col_keys"]
    nd = max(1, len(dims))
    nv = len(vals)
    fmts = [FORMATS.get(v["kind"]) for v in vals]
    groups = [" · ".join(x for x in ck if x) for ck in col_keys]
    width_cols = nd + (len(groups) + 1) * nv if col_keys else nd + nv
    for i in range(1, width_cols + 1):
        ws.column_dimensions[get_column_letter(i)].width = 28 if i <= nd else 16

    ws.append([_cell(ws, title or "Pivot hisobot", font=Font(bold=True, size=14))])
    ws.append([f"Tuzildi: {datetime.now(LOCAL_TZ).strftime('%Y-%m-%d %H:%M')} · {generated_by}"])
    for label, value in meta + cfg.describe():
        ws.append([_cell(ws, label, font=BOLD), value])
    ws.append([f"Zayavkalar: {result['orders_used']}" + (" · ekranda cheklangan, bu faylda hammasi" if result.get("truncated") else "")])
    ws.append([])

    dim_titles = [d["label"] for d in dims] or ["Jami"]
    if col_keys:
        top = [_cell(ws, "", fill=HEAD_FILL) for _ in range(nd)]
        for g in groups + ["Jami"]:
            top.append(_cell(ws, g, font=HEAD_FONT, fill=COL_FILL))
            top.extend(_cell(ws, "", fill=COL_FILL) for _ in range(nv - 1))
        ws.append(top)
        head = [_cell(ws, t, font=HEAD_FONT, fill=HEAD_FILL) for t in dim_titles]
        for _ in range(len(groups) + 1):
            head.extend(_cell(ws, v["label"], font=HEAD_FONT, fill=HEAD_FILL) for v in vals)
    else:
        head = [_cell(ws, t, font=HEAD_FONT, fill=HEAD_FILL) for t in dim_titles]
        head.extend(_cell(ws, v["label"], font=HEAD_FONT, fill=HEAD_FILL) for v in vals)
    ws.append(head)

    def numbers(row: dict, font=None, fill=None) -> list:
        out = []
        for chunk in (row.get("cells") or []) + [row["total"]]:
            for i, v in enumerate(chunk):
                out.append(_cell(ws, v, font=font, fill=fill, fmt=fmts[i]))
        return out

    for row in result["rows"]:
        if row["type"] == "subtotal":
            lvl = row.get("level", 0)
            labels = [_cell(ws, "", fill=SUB_FILL) for _ in range(nd)]
            labels[lvl] = _cell(ws, f"{row['keys'][-1]} — jami", font=BOLD, fill=SUB_FILL)
            ws.append(labels + numbers(row, BOLD, SUB_FILL))
        else:
            keys = row["keys"] or ["Jami"]
            ws.append([_cell(ws, k) for k in keys] + numbers(row))
    total_labels = [_cell(ws, "Umumiy jami", font=BOLD, fill=TOTAL_FILL)] + [_cell(ws, "", fill=TOTAL_FILL) for _ in range(nd - 1)]
    ws.append(total_labels + numbers(result["total"], BOLD, TOTAL_FILL))

    raw = wb.create_sheet("Ma’lumot")
    dim_list = list(DIMENSIONS.values())
    raw_titles = [d.label for d in dim_list] + ["Summa", "Og‘irlik, kg", "Tasdiq vaqti"]
    for i, t in enumerate(raw_titles, 1):
        raw.column_dimensions[get_column_letter(i)].width = max(12, min(30, len(t) + 6))
    raw.freeze_panes = "A2"
    raw.append([_cell(raw, t, font=HEAD_FONT, fill=HEAD_FILL) for t in raw_titles])
    for r in raw_rows[:raw_cap]:
        raw.append([_clean(d.get(r)) for d in dim_list] + [r.amount, r.weight, r.proof_at])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
